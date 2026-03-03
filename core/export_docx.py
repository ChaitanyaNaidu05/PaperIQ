import logging
from typing import Dict, List, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import io

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


def add_heading_with_color(doc, text: str, level: int = 1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(30, 58, 138)
    return heading


def add_table_with_style(doc, rows: int, cols: int):
    table = doc.add_table(rows=rows, cols=cols)
    table.style = 'Light Grid Accent 1'
    return table


def set_cell_background(cell, color_rgb):
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color_rgb)
    cell._element.get_or_add_tcPr().append(shading_elm)


def create_overview_section(doc, data: Dict, filename: str):
    add_heading_with_color(doc, 'Overview', level=1)
    
    scores = data.get('scores', {})
    composite = scores.get('Composite', 0)
    domain = data.get('domain', 'N/A')
    
    p = doc.add_paragraph()
    p.add_run('Document: ').bold = True
    p.add_run(filename)
    
    p = doc.add_paragraph()
    p.add_run('Composite Score: ').bold = True
    p.add_run(f'{composite:.1f}/100')
    
    p = doc.add_paragraph()
    p.add_run('Domain: ').bold = True
    p.add_run(domain)
    
    doc.add_paragraph()


def create_scores_table(doc, scores: Dict):
    add_heading_with_color(doc, 'Metric Scores', level=1)
    
    metrics = [(k, v) for k, v in scores.items() if k != 'Composite']
    
    table = add_table_with_style(doc, len(metrics) + 2, 2)
    
    header_cells = table.rows[0].cells
    header_cells[0].text = 'Metric'
    header_cells[1].text = 'Score (0-100)'
    
    for cell in header_cells:
        set_cell_background(cell, 'D9E9FF')
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    for idx, (metric, value) in enumerate(metrics, start=1):
        row_cells = table.rows[idx].cells
        row_cells[0].text = metric
        row_cells[1].text = f'{value:.1f}'
    
    composite_row = table.rows[-1].cells
    composite_row[0].text = 'Composite'
    composite_row[1].text = f'{scores.get("Composite", 0):.1f}'
    
    for cell in composite_row:
        set_cell_background(cell, 'E8F4FF')
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    doc.add_paragraph()


def create_statistics_section(doc, stats: Dict):
    add_heading_with_color(doc, 'Document Statistics', level=1)
    
    stat_items = [
        ('Word Count', f"{stats.get('word_count', 0):,}"),
        ('Sentence Count', f"{stats.get('sentence_count', 0):,}"),
        ('Avg Sentence Length', f"{stats.get('avg_sentence_len', 0):.1f} words"),
        ('Avg Word Length', f"{stats.get('avg_word_len', 0):.1f} chars"),
        ('Total Citations', str(stats.get('total_citations', 0))),
        ('Citation Density', f"{stats.get('citation_density_per_1k', 0):.2f} per 1k words"),
        ('Type-Token Ratio', f"{stats.get('type_token_ratio', 0):.4f}"),
        ('Complex Word Ratio', f"{stats.get('complex_word_ratio', 0):.2%}"),
    ]
    
    table = add_table_with_style(doc, len(stat_items) + 1, 2)
    
    header_cells = table.rows[0].cells
    header_cells[0].text = 'Statistic'
    header_cells[1].text = 'Value'
    
    for cell in header_cells:
        set_cell_background(cell, 'D9E9FF')
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    for idx, (label, value) in enumerate(stat_items, start=1):
        row_cells = table.rows[idx].cells
        row_cells[0].text = label
        row_cells[1].text = value
    
    doc.add_paragraph()


def create_keywords_section(doc, keywords: List):
    if not keywords:
        return
    
    add_heading_with_color(doc, 'Top Keywords', level=1)
    
    for kw, score in keywords[:15]:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f'{kw} ').bold = True
        p.add_run(f'(score: {score:.4f})')
    
    doc.add_paragraph()


def create_summary_section(doc, summary: str):
    if not summary:
        return
    
    add_heading_with_color(doc, 'Document Summary', level=1)
    
    p = doc.add_paragraph(summary)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    doc.add_paragraph()


def create_section_analysis(doc, section_scores: List[Dict]):
    if not section_scores:
        return
    
    add_heading_with_color(doc, 'Section-wise Analysis', level=1)
    
    table = add_table_with_style(doc, len(section_scores) + 1, 5)
    
    header_cells = table.rows[0].cells
    headers = ['Section', 'Score', 'Words', 'Clarity', 'Readability']
    for idx, header in enumerate(headers):
        header_cells[idx].text = header
        set_cell_background(header_cells[idx], 'D9E9FF')
        for paragraph in header_cells[idx].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    for idx, sec in enumerate(section_scores, start=1):
        row_cells = table.rows[idx].cells
        row_cells[0].text = sec.get('section', 'Unknown')[:40]
        row_cells[1].text = f"{sec.get('score', 0):.1f}"
        row_cells[2].text = str(sec.get('word_count', 0))
        row_cells[3].text = f"{sec.get('clarity_score', 0):.1f}"
        row_cells[4].text = f"{sec.get('readability', 0):.1f}"
    
    doc.add_paragraph()


def create_structure_section(doc, data: Dict):
    add_heading_with_color(doc, 'Structural Analysis', level=1)
    
    found = data.get('structural_found', [])
    missing = data.get('structural_missing', [])
    
    if found:
        p = doc.add_paragraph()
        p.add_run('Detected Sections:').bold = True
        for item in found:
            doc.add_paragraph(item, style='List Bullet')
    
    if missing:
        p = doc.add_paragraph()
        run = p.add_run('Missing Sections:')
        run.bold = True
        run.font.color.rgb = RGBColor(220, 38, 38)
        
        for item in missing:
            p = doc.add_paragraph(item, style='List Bullet')
            for run in p.runs:
                run.font.color.rgb = RGBColor(220, 38, 38)
    
    sentiment = data.get('sentiment', 0)
    p = doc.add_paragraph()
    p.add_run('Sentiment Analysis:').bold = True
    
    sentiment_text = f' {sentiment:.2f} '
    if sentiment > 0.1:
        sentiment_text += '(Positive)'
    elif sentiment < -0.1:
        sentiment_text += '(Negative)'
    else:
        sentiment_text += '(Neutral)'
    
    p.add_run(sentiment_text)
    
    doc.add_paragraph()


def create_advanced_analysis_section(doc, advanced: Dict):
    if not advanced:
        return
    
    add_heading_with_color(doc, 'Advanced Analysis', level=1)
    
    quality = advanced.get('quality_prediction', {})
    if quality:
        add_heading_with_color(doc, 'Quality Assessment', level=2)
        
        p = doc.add_paragraph()
        p.add_run('Quality Score: ').bold = True
        p.add_run(f"{quality.get('quality_score', 0):.1f}/100")
        
        p = doc.add_paragraph()
        p.add_run('Grade: ').bold = True
        p.add_run(quality.get('grade', 'N/A'))
        
        p = doc.add_paragraph()
        p.add_run('Percentile: ').bold = True
        p.add_run(f"{quality.get('percentile_estimate', 0)}th")
        
        strengths = quality.get('strengths', [])
        if strengths:
            p = doc.add_paragraph()
            p.add_run('Strengths:').bold = True
            for s in strengths[:5]:
                doc.add_paragraph(s, style='List Bullet')
        
        weaknesses = quality.get('weaknesses', [])
        if weaknesses:
            p = doc.add_paragraph()
            p.add_run('Areas for Improvement:').bold = True
            for w in weaknesses[:5]:
                doc.add_paragraph(w, style='List Bullet')
        
        doc.add_paragraph()
    
    reproducibility = advanced.get('reproducibility', {})
    if reproducibility:
        add_heading_with_color(doc, 'Reproducibility', level=2)
        
        rep_score = reproducibility.get('score', 0)
        rep_grade = reproducibility.get('grade', 'N/A')
        
        p = doc.add_paragraph()
        p.add_run('Score: ').bold = True
        p.add_run(f'{rep_score}/100')
        
        p = doc.add_paragraph()
        p.add_run('Grade: ').bold = True
        p.add_run(rep_grade)
        
        recommendations = reproducibility.get('recommendations', [])
        if recommendations:
            p = doc.add_paragraph()
            p.add_run('Recommendations:').bold = True
            for rec in recommendations[:5]:
                doc.add_paragraph(rec, style='List Bullet')
        
        doc.add_paragraph()
    
    writing = advanced.get('writing_quality', {})
    if writing:
        add_heading_with_color(doc, 'Writing Quality', level=2)
        
        p = doc.add_paragraph()
        p.add_run('Clarity Score: ').bold = True
        p.add_run(f"{writing.get('clarity_score', 0):.1f}/100")
        
        p = doc.add_paragraph()
        p.add_run('Avg Sentence Length: ').bold = True
        p.add_run(f"{writing.get('avg_sentence_length', 0):.1f} words")
        
        p = doc.add_paragraph()
        p.add_run('Passive Voice Ratio: ').bold = True
        p.add_run(f"{writing.get('passive_voice_ratio', 0)*100:.1f}%")
        
        doc.add_paragraph()


def create_docx_export(data: Dict, filename: str, section_scores: List[Dict] = None) -> bytes:
    try:
        doc = Document()
        
        title = doc.add_heading('PaperIQ Analysis Report', level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title.runs:
            run.font.color.rgb = RGBColor(30, 58, 138)
            run.font.size = Pt(24)
        
        subtitle = doc.add_paragraph('Generated by PaperIQ')
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
        
        create_overview_section(doc, data, filename)
        
        create_scores_table(doc, data.get('scores', {}))
        
        create_statistics_section(doc, data.get('stats', {}))
        
        keywords = data.get('keywords', [])
        if keywords:
            create_keywords_section(doc, keywords)
        
        summary = data.get('document_summary', '')
        if summary:
            create_summary_section(doc, summary)
        
        if section_scores:
            create_section_analysis(doc, section_scores)
        
        create_structure_section(doc, data)
        
        advanced = data.get('advanced', {})
        if advanced:
            create_advanced_analysis_section(doc, advanced)
        
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        return buffer.getvalue()
        
    except Exception as e:
        logger.error(f"Error generating DOCX export: {e}")
        return b''
